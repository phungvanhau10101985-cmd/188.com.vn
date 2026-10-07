"""Ghi file Excel báo cáo sau import sản phẩm — lý do bỏ qua từng dòng."""

from __future__ import annotations

import logging
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.services.import_excel_job_store import import_job_report_path

logger = logging.getLogger(__name__)

_HEADER_FONT = Font(bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")


def _as_text_list(value: Any) -> List[str]:
    if not isinstance(value, list):
        return []
    return [str(item).strip() for item in value if item is not None and str(item).strip()]


def _skip_rows(result: Dict[str, Any]) -> List[Dict[str, Any]]:
    raw = result.get("skipped_rows")
    rows: List[Dict[str, Any]] = []
    if isinstance(raw, list):
        for item in raw:
            if not isinstance(item, dict):
                continue
            reason = str(item.get("reason") or "").strip()
            if not reason:
                continue
            rows.append(
                {
                    "row": item.get("row") or "",
                    "source_id": str(item.get("source_id") or "").strip(),
                    "reason": reason,
                    "kept_product_id": str(item.get("kept_product_id") or "").strip(),
                }
            )
    if rows:
        return rows
    # Job cũ chỉ có câu mô tả — vẫn ghi vào báo cáo để không mất lý do.
    for line in _as_text_list(result.get("skipped")):
        rows.append(
            {
                "row": "",
                "source_id": "",
                "reason": line,
                "kept_product_id": "",
            }
        )
    return rows


def write_import_excel_job_report(job_id: str, result: Dict[str, Any]) -> bool:
    """
    Ghi ``{job_id}-report.xlsx`` khi import có dòng bỏ qua, lỗi hoặc cảnh báo.
    Trả False nếu không có gì để báo hoặc job_id không hợp lệ.
    """
    path = import_job_report_path(job_id)
    if path is None or not isinstance(result, dict):
        return False

    skips = _skip_rows(result)
    errors = _as_text_list(result.get("errors"))
    warnings = _as_text_list(result.get("warnings"))
    if not skips and not errors and not warnings:
        return False

    wb = Workbook()
    summary = wb.active
    summary.title = "Tóm tắt"
    summary.append(["Chỉ tiêu", "Giá trị"])
    summary["A1"].font = _HEADER_FONT
    summary["B1"].font = _HEADER_FONT
    for label, key in (
        ("Tạo mới", "created"),
        ("Cập nhật", "updated"),
        ("Xóa khỏi DB", "deleted"),
        ("Bỏ qua", "skipped_count"),
        ("Tổng dòng file", "total_processed"),
        ("Tỷ lệ không lỗi", "success_rate"),
    ):
        summary.append([label, result.get(key, "")])
    summary.append(["Số dòng lỗi", len(errors)])
    summary.append(["Số cảnh báo", len(warnings)])
    summary.append(
        [
            "Ghi chú",
            "Sheet «Bỏ qua» ghi lý do từng dòng không được ghi đè hoặc tạo mới.",
        ]
    )
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 88
    for cell in summary["B"]:
        cell.alignment = _WRAP

    skipped_sheet = wb.create_sheet("Bỏ qua")
    skipped_sheet.append(["Dòng", "Mã cột A", "Lý do bỏ qua", "Sản phẩm đang giữ"])
    for cell in skipped_sheet[1]:
        cell.font = _HEADER_FONT
    for item in skips:
        skipped_sheet.append(
            [
                item["row"],
                item["source_id"],
                item["reason"],
                item["kept_product_id"],
            ]
        )
    for col, width in ((1, 10), (2, 28), (3, 78), (4, 36)):
        skipped_sheet.column_dimensions[get_column_letter(col)].width = width
    for row in skipped_sheet.iter_rows(min_row=2, min_col=3, max_col=3):
        for cell in row:
            cell.alignment = _WRAP
    skipped_sheet.auto_filter.ref = skipped_sheet.dimensions
    skipped_sheet.freeze_panes = "A2"

    error_sheet = wb.create_sheet("Lỗi")
    error_sheet.append(["Chi tiết"])
    error_sheet["A1"].font = _HEADER_FONT
    for line in errors:
        error_sheet.append([line])
    error_sheet.column_dimensions["A"].width = 120
    for row in error_sheet.iter_rows(min_row=2, max_col=1):
        for cell in row:
            cell.alignment = _WRAP

    warning_sheet = wb.create_sheet("Cảnh báo")
    warning_sheet.append(["Chi tiết"])
    warning_sheet["A1"].font = _HEADER_FONT
    for line in warnings:
        warning_sheet.append([line])
    warning_sheet.column_dimensions["A"].width = 120

    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    tmp.replace(path)
    logger.info(
        "import excel report job=%s skips=%s errors=%s warnings=%s path=%s",
        job_id,
        len(skips),
        len(errors),
        len(warnings),
        path,
    )
    return True
